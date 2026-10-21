import csv
import datetime as dt
import io
from typing import Optional

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from .. import capacity, data_access
from ..config import get_settings
from ..database import get_db
from ..models import User
from ..security import get_current_user

router = APIRouter(prefix="/api/export", tags=["export"])


def _scope_keys(vcenter: Optional[str]):
    return None if (not vcenter or vcenter == "all") else [vcenter]


@router.get("/csv")
def export_csv(
    type: str = Query(default="hosts", pattern="^(hosts|vms|findings)$"),
    vcenter: Optional[str] = Query(default="all"),
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    settings = get_settings()
    policy = data_access.get_effective_policy(db, settings)
    host_rows, vm_rows, _ = data_access.get_latest_combined(db, _scope_keys(vcenter))

    buf = io.StringIO()
    if type == "hosts":
        rows = capacity.build_host_rows(host_rows, vm_rows, policy)
        fields = ["cluster", "host", "cores", "ramGB", "vcpuUsed", "ramUsedGB", "cpuPct", "ramPct", "overall"]
        writer = csv.writer(buf)
        writer.writerow(fields)
        for r in rows:
            writer.writerow([r.get("cluster"), r.get("host"), r.get("cores"), r.get("ramGB"),
                              r.get("vcpuUsed"), r.get("ramUsedGB"), r.get("cpuPct"), r.get("ramPct"),
                              r.get("overall", {}).get("text")])
    elif type == "vms":
        fields = ["cluster", "host", "vm", "powerState", "vcpu", "ramGB", "diskProvGB", "diskUsedGB", "datastore"]
        writer = csv.writer(buf)
        writer.writerow(fields)
        for v in vm_rows:
            writer.writerow([v.get(f) for f in fields])
    else:
        findings = capacity.build_compliance_findings(host_rows, vm_rows, policy)
        fields = ["severity", "type", "typeLabel", "entity", "cluster", "metric", "recommendation"]
        writer = csv.writer(buf)
        writer.writerow(fields)
        for f in findings:
            writer.writerow([f.get(k) for k in fields])

    buf.seek(0)
    filename = f"vco_{type}_{dt.datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.csv"
    return StreamingResponse(
        iter([buf.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ---------------------------------------------------------------------------
# Excel export of what the user currently sees in the Inventory tables.
# The browser already holds the filtered/sorted rows (global vCenter/cluster/
# host filter + table search/status/cluster filters), so it posts them here
# and gets a formatted .xlsx back — no need to replicate the UI filters.
# ---------------------------------------------------------------------------
from typing import Any, Dict, List  # noqa: E402

from fastapi import HTTPException  # noqa: E402
from pydantic import BaseModel, Field  # noqa: E402

MAX_XLSX_ROWS = 100_000


class XlsxColumn(BaseModel):
    key: str
    label: str
    fmt: Optional[str] = None  # "int" | "num1" | "pct" | None (text)


class XlsxSheet(BaseModel):
    name: str = Field(min_length=1, max_length=31)
    columns: List[XlsxColumn]
    rows: List[Dict[str, Any]]


class XlsxPayload(BaseModel):
    filename: str = "export"
    sheets: List[XlsxSheet]


def _safe_sheet_name(name: str) -> str:
    bad = '[]:*?/\\'
    return "".join(c for c in name if c not in bad)[:31] or "Sheet"


@router.post("/xlsx")
def export_xlsx(payload: XlsxPayload, user: User = Depends(get_current_user)):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:  # pragma: no cover
        raise HTTPException(status_code=501, detail="Máy chủ chưa cài openpyxl nên không xuất được Excel.")

    total_rows = sum(len(s.rows) for s in payload.sheets)
    if total_rows > MAX_XLSX_ROWS:
        raise HTTPException(status_code=413, detail=f"Quá nhiều dòng ({total_rows:,}); giới hạn {MAX_XLSX_ROWS:,}.")

    wb = Workbook()
    wb.remove(wb.active)
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="074B9B")
    number_formats = {"int": "#,##0", "num1": "#,##0.0", "pct": "0.0"}

    for sheet in payload.sheets or []:
        ws = wb.create_sheet(_safe_sheet_name(sheet.name))
        ws.append([c.label for c in sheet.columns])
        for cell in ws[1]:
            cell.font = head_font
            cell.fill = head_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for row in sheet.rows:
            values = []
            for c in sheet.columns:
                v = row.get(c.key)
                if isinstance(v, float) and (v != v or v in (float("inf"), float("-inf"))):
                    v = None  # NaN / Infinity → empty cell
                values.append(v)
            ws.append(values)
        widths = []
        for idx, c in enumerate(sheet.columns, 1):
            fmt = number_formats.get(c.fmt or "")
            longest = len(c.label)
            for r in range(2, ws.max_row + 1):
                cell = ws.cell(row=r, column=idx)
                if fmt and isinstance(cell.value, (int, float)):
                    cell.number_format = fmt
                longest = max(longest, len(str(cell.value)) if cell.value is not None else 0)
            widths.append(min(48, max(8, longest + 2)))
        for idx, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(idx)].width = w
        ws.freeze_panes = "A2"
        if sheet.columns:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(sheet.columns))}{max(1, ws.max_row)}"

    if not wb.sheetnames:
        wb.create_sheet("Sheet")
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    safe = "".join(ch for ch in payload.filename if ch.isalnum() or ch in "-_") or "export"
    filename = f"{safe}_{dt.datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return StreamingResponse(
        out,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
