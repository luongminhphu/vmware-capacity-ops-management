"""Host sizing API: calculate, import the Dev team's VM list, download a
template, and export the sizing result to Excel."""
import datetime as dt
import io
import unicodedata
from typing import Any, Dict, List, Optional

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session

from .. import report_sizing, sizing
from ..database import get_db
from ..models import ReportLogo, User
from ..security import get_current_user

router = APIRouter(prefix="/api/sizing", tags=["sizing"])

MAX_VMS = 5000
MAX_PROFILES = 8
MAX_UPLOAD = 5 * 1024 * 1024


class VmIn(BaseModel):
    name: str = ""
    qty: int = Field(default=1, ge=0, le=100000)
    vcpu: float = Field(default=0, ge=0, le=1024)
    ramGB: float = Field(default=0, ge=0, le=65536)
    diskGB: float = Field(default=0, ge=0)
    note: str = ""


class ProfileIn(BaseModel):
    name: str = ""
    sockets: int = Field(default=2, ge=1, le=8)
    coresPerSocket: int = Field(default=16, ge=1, le=256)
    ramGB: float = Field(default=512, ge=1, le=65536)
    diskGB: float = Field(default=0, ge=0)
    unitPrice: Optional[float] = Field(default=None, ge=0)


class ParamsIn(BaseModel):
    cpu_ratio: float = Field(default=3.0, gt=0, le=20)
    ht_factor: float = Field(default=2.0, ge=1, le=4)
    ram_ratio: float = Field(default=1.0, gt=0, le=4)
    max_util_pct: float = Field(default=80, ge=10, le=100)
    growth_pct: float = Field(default=20, ge=0, le=500)
    redundancy: int = Field(default=1, ge=0, le=4)
    min_hosts: int = Field(default=2, ge=1, le=64)
    reserve_cores: float = Field(default=2, ge=0, le=32)
    reserve_ram_gb: float = Field(default=16, ge=0, le=512)
    storage_mode: str = Field(default="san", pattern="^(san|local)$")
    storage_overhead_pct: float = Field(default=25, ge=0, le=300)
    storage_protection_factor: float = Field(default=1.0, ge=1, le=4)
    environment: str = Field(default="", max_length=120)


class ExistingClusterIn(BaseModel):
    """Absorption check computed by the UI from current vCenter data (optional, PDF only)."""
    name: str
    hostCount: int = 0
    cpuBefore: Optional[float] = None
    cpuAfter: Optional[float] = None
    ramBefore: Optional[float] = None
    ramAfter: Optional[float] = None
    n1CpuPct: Optional[float] = None
    n1RamPct: Optional[float] = None
    n1Ok: Optional[bool] = None
    cls: str = "na"
    totalCores: Optional[float] = None
    totalRamGB: Optional[float] = None
    usedVcpu: Optional[float] = None
    usedRamGB: Optional[float] = None
    poweredOnCount: Optional[int] = None
    cpuOcAfter: Optional[float] = None
    ramOcAfter: Optional[float] = None
    extraAvgVms: Optional[int] = None


class SizingPayload(BaseModel):
    vms: List[VmIn]
    params: ParamsIn = ParamsIn()
    profiles: List[ProfileIn]
    existing: Optional[List[ExistingClusterIn]] = None
    policyText: str = Field(default="", max_length=120)
    currency: str = Field(default="VND", max_length=8)
    preparedBy: str = Field(default="", max_length=80)
    # PDF content: 'new' = buy new servers (default) · 'reuse' = place the project on an existing cluster
    mode: str = Field(default="new", pattern="^(new|reuse)$")
    targetCluster: Optional[str] = Field(default=None, max_length=120)
    includeNewOption: bool = True      # reuse mode: append the buy-new comparison as option B
    includeExisting: bool = True       # new mode: append the existing-cluster absorption check


def _inputs(payload: SizingPayload):
    if len(payload.vms) > MAX_VMS:
        raise HTTPException(status_code=413, detail=f"Quá nhiều dòng VM ({len(payload.vms)}); giới hạn {MAX_VMS}.")
    if not payload.profiles:
        raise HTTPException(status_code=422, detail="Cần ít nhất một cấu hình host ứng viên.")
    if len(payload.profiles) > MAX_PROFILES:
        raise HTTPException(status_code=413, detail=f"Tối đa {MAX_PROFILES} cấu hình host ứng viên.")
    vms = [sizing.VmSpec(name=v.name or f"VM-{i+1:02d}", qty=v.qty, vcpu=v.vcpu, ram_gb=v.ramGB, disk_gb=v.diskGB, note=v.note)
           for i, v in enumerate(payload.vms)]
    profiles = [sizing.HostProfile(
        name=p.name or f"{p.sockets}×{p.coresPerSocket}C / {int(p.ramGB)} GB",
        sockets=p.sockets, cores_per_socket=p.coresPerSocket, ram_gb=p.ramGB, disk_gb=p.diskGB, unit_price=p.unitPrice,
    ) for p in payload.profiles]
    params = sizing.SizingParams(**payload.params.model_dump())
    return vms, params, profiles


def _build(payload: SizingPayload) -> Dict[str, Any]:
    vms, params, profiles = _inputs(payload)
    return sizing.build_sizing(vms, params, profiles)


@router.post("/pdf")
def sizing_pdf(payload: SizingPayload, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    vms, params, profiles = _inputs(payload)
    if not vms or not any(v.qty > 0 for v in vms):
        raise HTTPException(status_code=422, detail="Danh sách VM trống — không có gì để lập đề xuất.")
    logo_row = db.query(ReportLogo).filter(ReportLogo.id == 1).first()
    logo_data_url = f"data:{logo_row.mime_type};base64,{logo_row.data_base64}" if logo_row else None
    existing = [e.model_dump() for e in (payload.existing or [])][:20]
    try:
        pdf_bytes = report_sizing.build_sizing_pdf(vms, params, profiles, existing=existing, policy_txt=payload.policyText,
                                                   currency=payload.currency, logo_data_url=logo_data_url, prepared_by=payload.preparedBy,
                                                   mode=payload.mode, target_cluster=payload.targetCluster,
                                                   include_new=payload.includeNewOption, include_existing=payload.includeExisting)
    except Exception as exc:  # pragma: no cover - surfaced to the UI
        raise HTTPException(status_code=500, detail=f"Không thể tạo PDF: {exc}")
    label = _ascii_slug(params.environment, fallback="")
    kind = "tan_dung_hien_huu" if payload.mode == "reuse" else "trang_bi_moi"
    filename = f"VCO_De_xuat_cau_hinh_may_chu_{(label + '_') if label else ''}{kind}_{dt.datetime.now().strftime('%Y%m%d_%H%M')}.pdf"
    return StreamingResponse(iter([pdf_bytes]), media_type="application/pdf", headers=_attachment_headers(filename))


@router.get("/defaults")
def sizing_defaults(user: User = Depends(get_current_user)):
    return {"params": sizing.SizingParams().__dict__, "profiles": sizing.DEFAULT_PROFILES, "sampleVms": sizing.TEMPLATE_ROWS}


@router.post("/calc")
def sizing_calc(payload: SizingPayload, user: User = Depends(get_current_user)):
    return _build(payload)


@router.post("/import")
async def sizing_import(file: UploadFile = File(...), user: User = Depends(get_current_user)):
    content = await file.read()
    if len(content) > MAX_UPLOAD:
        raise HTTPException(status_code=413, detail="File vượt quá 5 MB.")
    try:
        parsed = sizing.parse_vm_file(file.filename or "", content)
    except Exception as exc:  # noqa: BLE001 — surface a readable message to the UI
        raise HTTPException(status_code=400, detail=f"Không đọc được file: {exc}")
    return parsed


def _ascii_slug(text: Optional[str], fallback: str = "may_chu", limit: int = 40) -> str:
    """ASCII-only file-name fragment. HTTP headers are latin-1, so Vietnamese
    project names ("Dự án Đầu tư / CRM 2026") must be transliterated, not just
    filtered with str.isalnum() (which accepts Unicode letters)."""
    text = (text or "").replace("đ", "d").replace("Đ", "D")
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
    slug = "".join(ch if (ch.isalnum() or ch in "-_") else "_" for ch in text.strip()).strip("_")
    while "__" in slug:
        slug = slug.replace("__", "_")
    return slug[:limit].strip("_") or fallback


def _attachment_headers(filename: str) -> Dict[str, str]:
    safe = filename.encode("ascii", "ignore").decode("ascii").replace('"', "")
    return {"Content-Disposition": f'attachment; filename="{safe}"'}


def _xlsx_response(wb, base: str) -> StreamingResponse:
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    filename = f"{base}_{dt.datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return StreamingResponse(
        out,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=_attachment_headers(filename),
    )


def _style_header(ws, row: int = 1):
    from openpyxl.styles import Alignment, Font, PatternFill
    for cell in ws[row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="074B9B")
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _autowidth(ws, min_w: int = 8, max_w: int = 60):
    from openpyxl.utils import get_column_letter
    for idx in range(1, ws.max_column + 1):
        longest = 0
        for r in range(1, ws.max_row + 1):
            v = ws.cell(row=r, column=idx).value
            longest = max(longest, len(str(v)) if v is not None else 0)
        ws.column_dimensions[get_column_letter(idx)].width = min(max_w, max(min_w, longest + 2))


@router.get("/template")
def sizing_template(user: User = Depends(get_current_user)):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "VM list"
    ws.append(["Tên VM", "Số lượng", "vCPU", "RAM (GB)", "Disk (GB)", "Ghi chú"])
    _style_header(ws)
    for r in sizing.TEMPLATE_ROWS:
        ws.append([r["name"], r["qty"], r["vcpu"], r["ramGB"], r["diskGB"], r["note"]])
    ws.freeze_panes = "A2"
    _autowidth(ws)
    guide = wb.create_sheet("Hướng dẫn")
    for line in [
        "Điền mỗi dòng một loại VM; cột 'Số lượng' cho phép khai nhiều VM cùng cấu hình.",
        "Đơn vị RAM và Disk là GB (có thể ghi '2 TB', '512 GB' — hệ thống tự quy đổi).",
        "Tên cột có thể viết tiếng Việt hoặc tiếng Anh (VM name / Qty / vCPU / RAM / Disk / Note).",
        "Lưu file .xlsx hoặc .csv rồi dùng nút 'Import Excel/CSV' trong mục Sizing máy chủ.",
    ]:
        guide.append([line])
    guide.column_dimensions["A"].width = 110
    return _xlsx_response(wb, "VCO_Sizing_mau_danh_sach_VM")


@router.post("/export")
def sizing_export(payload: SizingPayload, user: User = Depends(get_current_user)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    result = _build(payload)
    demand = result["demand"]
    params = result["params"]
    wb = Workbook()

    # Sheet 1 — summary & recommendation
    ws = wb.active
    ws.title = "Tóm tắt"
    ws.append(["Sizing máy chủ vật lý — " + (params.get("environment") or "dự án mới")])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Ngày: {dt.date.today().strftime('%d/%m/%Y')}"])
    ws.append([])
    ws.append(["Đề xuất", result["recommendation"]["headline"]])
    ws.append(["Chi tiết", result["recommendation"]["detail"]])
    ws.append([])
    ws.append(["Nhu cầu từ team Dev", "Giá trị", "Sau dự phòng tăng trưởng"])
    _style_header(ws, ws.max_row)
    ws.append(["Số VM", demand["vmCount"], demand["vmCount"]])
    ws.append(["vCPU", demand["vcpu"], round(demand["vcpuPlanned"], 1)])
    ws.append(["RAM (GB)", demand["ramGB"], round(demand["ramPlannedGB"], 1)])
    ws.append(["Disk (GB)", demand["diskGB"], round(demand["diskPlannedGB"], 1)])
    ws.append(["Dung lượng lưu trữ cần cấp (GB, gồm overhead)", "", round(demand["storageNeededGB"], 1)])
    ws.append([])
    ws.append(["Giả định"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    for line in result["assumptions"]:
        ws.append([line])
    _autowidth(ws, max_w=90)

    # Sheet 2 — candidate comparison
    ws2 = wb.create_sheet("So sánh cấu hình")
    headers = ["Cấu hình host", "Sockets × cores", "RAM/host (GB)", "Số host theo CPU", "Số host theo RAM",
               "Dự phòng N+x", "Tổng số host", "Ràng buộc", "Tổng cores", "Tổng RAM (GB)",
               "Overcommit CPU sau", "Overcommit RAM sau", "CPU % khi mất host", "RAM % khi mất host",
               "Đạt N+x", "Thêm được ~VM", "Đơn giá", "Ngân sách ước tính", "Đề xuất"]
    ws2.append(headers)
    _style_header(ws2)
    for r in result["results"]:
        p = r["profile"]
        ws2.append([
            p["name"], f"{p['sockets']}×{p['coresPerSocket']}", p["ramGB"], r["hostsForCpu"], r["hostsForRam"],
            r["redundancy"], r["hostsTotal"], r["binding"], r["totals"]["cores"], r["totals"]["ramGB"],
            round(r["after"]["cpuOvercommit"], 2) if r["after"]["cpuOvercommit"] is not None else None,
            round(r["after"]["ramOvercommit"], 2) if r["after"]["ramOvercommit"] is not None else None,
            round(r["failure"]["cpuPct"], 1) if r["failure"]["cpuPct"] is not None else None,
            round(r["failure"]["ramPct"], 1) if r["failure"]["ramPct"] is not None else None,
            "Đạt" if r["failure"]["ok"] else "Chưa đạt",
            r["spare"]["extraAvgVms"], p["unitPrice"], r["budget"], "✓" if r["recommended"] else "",
        ])
        if r["recommended"]:
            for cell in ws2[ws2.max_row]:
                cell.fill = PatternFill("solid", fgColor="E6F4EA")
                cell.font = Font(bold=True)
    ws2.freeze_panes = "B2"
    _autowidth(ws2)

    # Sheet 3 — VM list as submitted
    ws3 = wb.create_sheet("Danh sách VM")
    ws3.append(["Tên VM", "Số lượng", "vCPU", "RAM (GB)", "Disk (GB)", "Tổng vCPU", "Tổng RAM (GB)", "Tổng Disk (GB)", "Ghi chú"])
    _style_header(ws3)
    for v in payload.vms:
        ws3.append([v.name, v.qty, v.vcpu, v.ramGB, v.diskGB, v.qty * v.vcpu, v.qty * v.ramGB, v.qty * v.diskGB, v.note])
    ws3.append(["TỔNG", demand["vmCount"], "", "", "", demand["vcpu"], demand["ramGB"], demand["diskGB"], ""])
    for cell in ws3[ws3.max_row]:
        cell.font = Font(bold=True)
    ws3.freeze_panes = "A2"
    _autowidth(ws3)

    base = "VCO_Sizing" + (("_" + _ascii_slug(params.get("environment"), fallback="")) if (params.get("environment") or "").strip() else "")
    return _xlsx_response(wb, base)
